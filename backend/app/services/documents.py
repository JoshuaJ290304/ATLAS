from __future__ import annotations

import asyncio
import hashlib
import re
import shutil
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from uuid import UUID, uuid4

import fitz
from docx import Document as DocxDocument
from fastapi import HTTPException, UploadFile, status

from app.core.config import settings

CHUNK_SIZE = 1024 * 1024
CONTROL_CHARACTERS = re.compile(r"[\x00-\x1f\x7f]")
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True)
class StoredUpload:
    original_filename: str
    stored_filename: str
    storage_path: Path
    media_type: str
    extension: str
    size_bytes: int
    checksum_sha256: str


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    page_count: int | None


def clean_filename(filename: str | None) -> str:
    cleaned = Path((filename or "").replace("\\", "/")).name.strip()
    cleaned = CONTROL_CHARACTERS.sub("", cleaned)
    if not cleaned or cleaned in {".", ".."}:
        raise HTTPException(status_code=400, detail="A valid filename is required")
    if len(cleaned) > 255:
        raise HTTPException(status_code=400, detail="Filename must be 255 characters or fewer")
    return cleaned


def allowed_extensions() -> set[str]:
    return {
        extension.strip().lower()
        for extension in settings.ALLOWED_DOCUMENT_EXTENSIONS.split(",")
        if extension.strip()
    }


def validate_extension(filename: str, extensions: set[str] | None = None) -> str:
    configured = extensions or allowed_extensions()
    extension = Path(filename).suffix.lower()
    if extension not in configured:
        allowed = ", ".join(sorted(configured))
        raise HTTPException(
            status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
            detail=f"Unsupported file type. Allowed extensions: {allowed}",
        )
    return extension


def upload_root() -> Path:
    root = Path(settings.UPLOAD_ROOT).expanduser().resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def project_upload_directory(project_id: UUID) -> Path:
    directory = upload_root() / str(project_id)
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def resolve_stored_path(storage_path: str) -> Path:
    root = upload_root()
    candidate = Path(storage_path).resolve()
    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise HTTPException(status_code=500, detail="Invalid document storage path") from exc
    return candidate


def _validate_docx_archive(
    path: Path,
    *,
    max_uncompressed_mb: int,
    max_compression_ratio: int,
) -> None:
    if not zipfile.is_zipfile(path):
        raise HTTPException(status_code=422, detail="The file is not a valid DOCX document")

    max_uncompressed_bytes = max_uncompressed_mb * 1024 * 1024
    total_uncompressed = 0
    total_compressed = 0

    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise HTTPException(
                    status_code=422,
                    detail="The file is not a valid DOCX document",
                )

            for member in archive.infolist():
                member_path = PurePosixPath(member.filename)
                if member.flag_bits & 0x1:
                    raise HTTPException(
                        status_code=422,
                        detail="Encrypted DOCX files are not supported",
                    )
                if member_path.is_absolute() or ".." in member_path.parts:
                    raise HTTPException(
                        status_code=422,
                        detail="The DOCX archive contains an unsafe path",
                    )
                total_uncompressed += member.file_size
                total_compressed += member.compress_size
                if total_uncompressed > max_uncompressed_bytes:
                    raise HTTPException(
                        status_code=422,
                        detail=("The DOCX expands beyond the configured uncompressed-size limit"),
                    )

            if total_uncompressed and total_compressed == 0:
                raise HTTPException(
                    status_code=422,
                    detail="The DOCX archive has an unsafe compression profile",
                )
            if total_compressed:
                compression_ratio = total_uncompressed / total_compressed
                if compression_ratio > max_compression_ratio:
                    raise HTTPException(
                        status_code=422,
                        detail="The DOCX archive has an unsafe compression ratio",
                    )

            if archive.testzip() is not None:
                raise HTTPException(status_code=422, detail="The DOCX archive is corrupted")
    except zipfile.BadZipFile as exc:
        raise HTTPException(
            status_code=422,
            detail="The file is not a valid DOCX document",
        ) from exc


def validate_file_signature(
    path: Path,
    extension: str,
    *,
    max_docx_uncompressed_mb: int | None = None,
    max_docx_compression_ratio: int | None = None,
) -> None:
    if extension == ".pdf":
        with path.open("rb") as source:
            if source.read(5) != b"%PDF-":
                raise HTTPException(status_code=422, detail="The file is not a valid PDF")
        return

    if extension == ".docx":
        _validate_docx_archive(
            path,
            max_uncompressed_mb=(max_docx_uncompressed_mb or settings.MAX_DOCX_UNCOMPRESSED_MB),
            max_compression_ratio=(
                max_docx_compression_ratio or settings.MAX_DOCX_COMPRESSION_RATIO
            ),
        )
        return

    raise HTTPException(status_code=415, detail="Unsupported document type")


def remove_project_upload_directory(project_id: UUID) -> None:
    directory = upload_root() / str(project_id)
    if directory.is_dir():
        shutil.rmtree(directory)


async def store_upload(
    upload: UploadFile,
    project_id: UUID,
    *,
    max_size_mb: int | None = None,
    extensions: set[str] | None = None,
) -> StoredUpload:
    original_filename = clean_filename(upload.filename)
    extension = validate_extension(original_filename, extensions)
    stored_filename = f"{uuid4().hex}{extension}"
    destination = project_upload_directory(project_id) / stored_filename
    effective_max_mb = max_size_mb or settings.MAX_UPLOAD_SIZE_MB
    max_bytes = effective_max_mb * 1024 * 1024
    size_bytes = 0
    checksum = hashlib.sha256()

    try:
        with destination.open("wb") as output:
            while chunk := await upload.read(CHUNK_SIZE):
                size_bytes += len(chunk)
                if size_bytes > max_bytes:
                    raise HTTPException(
                        status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                        detail=f"File exceeds the {effective_max_mb} MB limit",
                    )
                checksum.update(chunk)
                output.write(chunk)
    except Exception:
        destination.unlink(missing_ok=True)
        raise
    finally:
        await upload.close()

    if size_bytes == 0:
        destination.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="The uploaded file is empty")

    return StoredUpload(
        original_filename=original_filename,
        stored_filename=stored_filename,
        storage_path=destination,
        media_type=MEDIA_TYPES[extension],
        extension=extension,
        size_bytes=size_bytes,
        checksum_sha256=checksum.hexdigest(),
    )


def _extract_pdf(path: Path) -> ExtractionResult:
    with fitz.open(path) as pdf:
        text = "\n\n".join(page.get_text("text").strip() for page in pdf).strip()
        return ExtractionResult(text=text, page_count=pdf.page_count)


def _extract_docx(path: Path) -> ExtractionResult:
    document = DocxDocument(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(value for value in cells if value)
            if line:
                parts.append(line)

    return ExtractionResult(text="\n".join(parts).strip(), page_count=None)


def _validate_extraction_result(result: ExtractionResult) -> ExtractionResult:
    if len(result.text) > settings.MAX_EXTRACTED_TEXT_CHARS:
        raise ValueError("Extracted text exceeds the configured character limit")
    return result


async def extract_document(path: Path, extension: str) -> ExtractionResult:
    if extension == ".pdf":
        task = asyncio.to_thread(_extract_pdf, path)
    elif extension == ".docx":
        task = asyncio.to_thread(_extract_docx, path)
    else:
        raise ValueError(f"Unsupported extension: {extension}")

    result = await asyncio.wait_for(
        task,
        timeout=settings.DOCUMENT_EXTRACTION_TIMEOUT_SECONDS,
    )
    return _validate_extraction_result(result)
