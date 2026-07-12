from pathlib import Path

import fitz
import pytest
from docx import Document as DocxDocument
from fastapi import HTTPException

from app.services.documents import (
    _extract_docx,
    _extract_pdf,
    clean_filename,
    validate_file_signature,
)


def test_clean_filename_removes_client_path() -> None:
    assert clean_filename(r"C:\\fakepath\\requirements.pdf") == "requirements.pdf"
    assert clean_filename("folder/requirements.pdf") == "requirements.pdf"


def test_extract_pdf(tmp_path: Path) -> None:
    path = tmp_path / "sample.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "REQ-001 Users shall log in securely.")
    pdf.save(path)
    pdf.close()

    result = _extract_pdf(path)

    assert result.page_count == 1
    assert "REQ-001" in result.text


def test_extract_docx_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("REQ-002 The system shall generate reports.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Priority"
    table.cell(0, 1).text = "High"
    document.save(path)

    result = _extract_docx(path)

    assert result.page_count is None
    assert "REQ-002" in result.text
    assert "Priority | High" in result.text


def test_validate_file_signature_rejects_fake_pdf(tmp_path: Path) -> None:
    path = tmp_path / "fake.pdf"
    path.write_text("not a pdf")

    with pytest.raises(HTTPException) as error:
        validate_file_signature(path, ".pdf")

    assert error.value.status_code == 422
